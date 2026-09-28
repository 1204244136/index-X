from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

TOOLS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOLS))

import openpyxl

from read_xlsx import (
    _col_letter_to_index,
    _format_cell_value,
    _parse_col_selection,
    format_csv,
    format_json,
    format_markdown,
    format_tsv,
    list_sheets,
    load_sheet_data,
    main,
)


class ReadXlsxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.test_file = Path(cls.temp_dir.name) / "sample.xlsx"

        # 创建一个多工作表的测试 xlsx
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Terms"
        ws1.append(["ID", "JP_Name", "CN_Name", "Category", "Notes"])
        ws1.append([1, "上条当麻", "上条当麻", "人物", "刺猬头\n男主角"])
        ws1.append([2, "御坂美琴", "御坂美琴", "人物", "第三位|超电磁炮"])
        ws1.append([3, "幻想殺し", "幻想杀手", "能力", "Imagine Breaker"])
        ws1.append([4, "一方通行", "一方通行", "人物", "第一位"])

        ws2 = wb.create_sheet(title="EmptySheet")

        ws3 = wb.create_sheet(title="Numbers")
        ws3.append(["Code", "Value", "Ratio"])
        ws3.append(["A01", 100.0, 0.25])
        ws3.append(["A02", 200, 0.5])

        wb.save(cls.test_file)
        wb.close()

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()

    def test_col_letter_to_index(self):
        self.assertEqual(_col_letter_to_index("A"), 0)
        self.assertEqual(_col_letter_to_index("B"), 1)
        self.assertEqual(_col_letter_to_index("Z"), 25)
        self.assertEqual(_col_letter_to_index("AA"), 26)
        self.assertEqual(_col_letter_to_index("AB"), 27)

    def test_format_cell_value(self):
        self.assertEqual(_format_cell_value(None), "")
        self.assertEqual(_format_cell_value(100.0), "100")
        self.assertEqual(_format_cell_value(3.1415), "3.1415")
        self.assertEqual(_format_cell_value("  hello  "), "hello")
        self.assertEqual(_format_cell_value(True), "True")

    def test_parse_col_selection(self):
        headers = ["ID", "JP_Name", "CN_Name", "Category", "Notes"]
        # 按名称
        res = _parse_col_selection("JP_Name,CN_Name", headers, 5)
        self.assertEqual(res, [1, 2])
        # 按字母
        res = _parse_col_selection("A,C", headers, 5)
        self.assertEqual(res, [0, 2])
        # 按 1-based 数字
        res = _parse_col_selection("2,4", headers, 5)
        self.assertEqual(res, [1, 3])
        # 空选择返回全部
        res = _parse_col_selection("", headers, 5)
        self.assertEqual(res, [0, 1, 2, 3, 4])

    def test_list_sheets(self):
        sheets = list_sheets(self.test_file)
        names = [s["name"] for s in sheets]
        self.assertEqual(names, ["Terms", "EmptySheet", "Numbers"])

    def test_load_sheet_data_default(self):
        sheet_name, headers, rows, total = load_sheet_data(self.test_file)
        self.assertEqual(sheet_name, "Terms")
        self.assertEqual(headers, ["ID", "JP_Name", "CN_Name", "Category", "Notes"])
        self.assertEqual(total, 4)
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0][1], "上条当麻")

    def test_load_sheet_data_head_and_tail(self):
        # head
        _, _, rows_head, total = load_sheet_data(self.test_file, head=2)
        self.assertEqual(total, 4)
        self.assertEqual(len(rows_head), 2)
        self.assertEqual(rows_head[0][1], "上条当麻")
        self.assertEqual(rows_head[1][1], "御坂美琴")

        # tail
        _, _, rows_tail, total = load_sheet_data(self.test_file, tail=2)
        self.assertEqual(total, 4)
        self.assertEqual(len(rows_tail), 2)
        self.assertEqual(rows_tail[0][1], "幻想殺し")
        self.assertEqual(rows_tail[1][1], "一方通行")

    def test_load_sheet_data_search(self):
        # 搜索“幻想”
        _, _, rows, total = load_sheet_data(self.test_file, search="幻想")
        self.assertEqual(total, 1)
        self.assertEqual(rows[0][1], "幻想殺し")
        self.assertEqual(rows[0][2], "幻想杀手")

        # 搜索不区分大小写
        _, _, rows_case, total_case = load_sheet_data(self.test_file, search="imagine")
        self.assertEqual(total_case, 1)
        self.assertEqual(rows_case[0][1], "幻想殺し")

    def test_load_sheet_data_cols(self):
        _, headers, rows, _ = load_sheet_data(self.test_file, cols="JP_Name,Notes")
        self.assertEqual(headers, ["JP_Name", "Notes"])
        self.assertEqual(len(rows[0]), 2)
        self.assertEqual(rows[0][0], "上条当麻")

    def test_format_markdown_sanitization(self):
        # 测试换行符转为 <br> 以及竖线转义
        headers = ["Col1", "Col2"]
        rows = [["Line1\nLine2", "Val|With|Pipe"]]
        md = format_markdown(headers, rows)
        self.assertIn("Line1<br>Line2", md)
        self.assertIn(r"Val\|With\|Pipe", md)
        self.assertTrue(md.startswith("| Col1"))

    def test_format_json_modes(self):
        headers = ["A", "B"]
        rows = [["1", "2"]]
        # records
        json_rec = format_json(headers, rows, mode="records")
        parsed_rec = json.loads(json_rec)
        self.assertEqual(parsed_rec, [{"A": "1", "B": "2"}])

        # rows
        json_rows = format_json(headers, rows, mode="rows")
        parsed_rows = json.loads(json_rows)
        self.assertEqual(parsed_rows, {"headers": ["A", "B"], "rows": [["1", "2"]]})

    def test_cli_execution(self):
        # 测试 main CLI 函数
        ret = main([str(self.test_file), "--list-sheets"])
        self.assertEqual(ret, 0)

        ret = main([str(self.test_file), "--head", "2", "--format", "tsv"])
        self.assertEqual(ret, 0)

        ret = main([str(self.test_file), "-q", "御坂美琴", "--format", "json"])
        self.assertEqual(ret, 0)


if __name__ == "__main__":
    unittest.main()
