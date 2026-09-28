#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel (.xlsx) 文件解析与预览工具。

用于快速查看、过滤、提取并转换 .xlsx 文件内容为模型友好或易读的格式（Markdown、TSV、CSV、JSON），
避免在对话中重复临时编写一次性解析脚本。

特性：
- 格式转换：支持转换为 Markdown 表格、TSV、CSV、JSON（records 或 rows）。
- 上下文保护：支持 --head、--tail、--limit、--offset 切片，防止大表格撑爆模型上下文。
- 列选择：支持按列名或列字母/数字（如 "A,B,D" 或 "1,2,4" 或 "Base_Ori,Base_Trans"）提取指定列。
- 行搜索：支持 --search / -q 关键词过滤，直接定位包含指定关键词的行。
- 工作表管理：支持 --list-sheets 查看所有 Sheet 及其行列规模，支持按名称或序号指定 Sheet。
- UTF-8 安全：强制标准输出采用 UTF-8 编码，消除 Windows 控制台中文乱码。

常用示例：
    # 1. 查看包含哪些工作表
    python tools/read_xlsx.py data.xlsx --list-sheets

    # 2. 预览默认工作表前 20 行（Markdown 格式）
    python tools/read_xlsx.py data.xlsx --head 20

    # 3. 指定工作表、仅提取指定列并输出为 Markdown
    python tools/read_xlsx.py data.xlsx --sheet data --cols "Base_Ori,Base_Trans,Type" --head 15

    # 4. 按关键词搜索包含“表层融解”的行，输出为 Markdown
    python tools/read_xlsx.py data.xlsx -q "表层融解"

    # 5. 导出为 TSV 或 JSON 文件
    python tools/read_xlsx.py data.xlsx --format json -o output.json
    python tools/read_xlsx.py data.xlsx --format tsv --head 100 -o sample.tsv
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import sys
from pathlib import Path
from typing import Any, Iterable, List, Optional, Tuple, Union

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore


def _ensure_utf8_stdout() -> None:
    """确保在 Windows 终端中输出 UTF-8 文本而不发生编码报错。"""
    try:
        if sys.stdout and hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        if sys.stderr and hasattr(sys.stderr, "reconfigure"):
            sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def _format_cell_value(val: Any) -> str:
    """将单元格值格式化为可打印字符串。"""
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:g}"
    if isinstance(val, (int, bool)):
        return str(val)
    return str(val).strip()


def list_sheets(file_path: str | Path) -> list[dict[str, Any]]:
    """列出工作簿中的所有工作表及其基本信息。"""
    if openpyxl is None:
        raise RuntimeError("未检测到 openpyxl 库，请先安装：pip install openpyxl")

    path = Path(file_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"文件不存在: {path}")

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    results = []
    try:
        for idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            # read_only 模式下 max_row / max_column 可能是预估或准确的
            results.append({
                "index": idx,
                "name": sheet_name,
                "max_row": getattr(ws, "max_row", None),
                "max_column": getattr(ws, "max_column", None),
            })
    finally:
        wb.close()
    return results


def _col_letter_to_index(col_str: str) -> int:
    """将 Excel 列字母（如 A, B, Z, AA）转换为 0-based 索引。"""
    col_str = col_str.strip().upper()
    idx = 0
    for char in col_str:
        if not ("A" <= char <= "Z"):
            raise ValueError(f"无效的列字母: {col_str}")
        idx = idx * 26 + (ord(char) - ord("A") + 1)
    return idx - 1


def _parse_col_selection(
    cols_spec: str,
    headers: list[str],
    total_cols: int,
) -> list[int]:
    """根据列名、字母或索引解析选中的列索引列表（0-based）。"""
    if not cols_spec or not cols_spec.strip():
        return list(range(total_cols))

    raw_items = [c.strip() for c in cols_spec.split(",") if c.strip()]
    selected: list[int] = []

    header_to_idx = {h.strip(): i for i, h in enumerate(headers) if h.strip()}

    for item in raw_items:
        # 1. 尝试直接匹配表头名称
        if item in header_to_idx:
            selected.append(header_to_idx[item])
            continue

        # 2. 尝试解析为 1-based 数字索引
        if item.isdigit():
            val = int(item) - 1
            if 0 <= val < total_cols:
                selected.append(val)
            continue

        # 3. 尝试解析为列字母 (A, B, AA...)
        if re.fullmatch(r"[A-Za-z]+", item):
            try:
                c_idx = _col_letter_to_index(item)
                if 0 <= c_idx < total_cols:
                    selected.append(c_idx)
                    continue
            except ValueError:
                pass

        # 4. 若不区分大小写匹配表头
        matched = False
        item_lower = item.lower()
        for h, idx in header_to_idx.items():
            if h.lower() == item_lower:
                selected.append(idx)
                matched = True
                break
        if matched:
            continue

        # 若未找到匹配列，向用户发出友好告警并忽略
        sys.stderr.write(f"[警告] 未找到指定列: '{item}'\n")

    # 去重并保持顺序
    seen = set()
    result = []
    for idx in selected:
        if idx not in seen and 0 <= idx < total_cols:
            seen.add(idx)
            result.append(idx)

    return result if result else list(range(total_cols))


def load_sheet_data(
    file_path: str | Path,
    sheet: str | int | None = None,
    header_row: int = 1,
    cols: str | None = None,
    search: str | None = None,
    search_ignore_case: bool = True,
    offset: int = 0,
    limit: int | None = None,
    head: int | None = None,
    tail: int | None = None,
) -> tuple[str, list[str], list[list[str]], int]:
    """加载并过滤 Excel 工作表数据。

    返回:
        (actual_sheet_name, headers, filtered_rows, total_data_rows)
    """
    if openpyxl is None:
        raise RuntimeError("未检测到 openpyxl 库，请先安装：pip install openpyxl")

    path = Path(file_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"文件不存在: {path}")

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheetnames = wb.sheetnames
        if not sheetnames:
            raise ValueError(f"工作簿为空，没有工作表: {path}")

        if sheet is None:
            actual_sheet = sheetnames[0]
        elif isinstance(sheet, int) or (isinstance(sheet, str) and sheet.isdigit()):
            s_idx = int(sheet)
            if 0 <= s_idx < len(sheetnames):
                actual_sheet = sheetnames[s_idx]
            else:
                raise ValueError(f"工作表序号超出范围: {sheet} (有效范围 0..{len(sheetnames)-1})")
        else:
            if sheet in sheetnames:
                actual_sheet = sheet
            else:
                # 尝试不区分大小写匹配
                matched = [s for s in sheetnames if s.lower() == str(sheet).lower()]
                if matched:
                    actual_sheet = matched[0]
                else:
                    raise ValueError(f"未找到工作表: '{sheet}'，可用表: {sheetnames}")

        ws = wb[actual_sheet]

        # 读取所有行（使用 read_only 生成器，内存占用低）
        raw_rows_iter = ws.iter_rows(values_only=True)

        current_row_idx = 0
        headers_raw: list[Any] = []
        data_rows: list[list[str]] = []

        # 处理表头
        if header_row > 0:
            for row in raw_rows_iter:
                current_row_idx += 1
                if current_row_idx == header_row:
                    headers_raw = list(row)
                    break
        else:
            # 无表头
            pass

        # 确定列总数
        for row in raw_rows_iter:
            current_row_idx += 1
            formatted = [_format_cell_value(c) for c in row]
            # 过滤全空行
            if any(formatted):
                data_rows.append(formatted)

    finally:
        wb.close()

    total_cols = 0
    if headers_raw:
        total_cols = max(total_cols, len(headers_raw))
    for r in data_rows:
        total_cols = max(total_cols, len(r))

    # 规范化 headers
    headers: list[str] = []
    if header_row > 0 and headers_raw:
        for idx in range(total_cols):
            val = headers_raw[idx] if idx < len(headers_raw) else ""
            h_str = _format_cell_value(val)
            headers.append(h_str if h_str else f"Col_{idx + 1}")
    else:
        headers = [f"Col_{idx + 1}" for idx in range(total_cols)]

    # 补齐 data_rows 的长度
    normalized_rows: list[list[str]] = []
    for r in data_rows:
        if len(r) < total_cols:
            r = r + [""] * (total_cols - len(r))
        normalized_rows.append(r)

    # 列筛选
    selected_col_indices = _parse_col_selection(cols or "", headers, total_cols)
    final_headers = [headers[i] for i in selected_col_indices]
    col_filtered_rows = [[r[i] for i in selected_col_indices] for r in normalized_rows]

    # 行关键词搜索筛选
    if search:
        search_target = search.lower() if search_ignore_case else search
        matched_rows: list[list[str]] = []
        for r in col_filtered_rows:
            matched = False
            for cell in r:
                c_val = cell.lower() if search_ignore_case else cell
                if search_target in c_val:
                    matched = True
                    break
            if matched:
                matched_rows.append(r)
        processed_rows = matched_rows
    else:
        processed_rows = col_filtered_rows

    total_matched_rows = len(processed_rows)

    # 切片处理: head, tail, offset/limit
    if head is not None and head >= 0:
        processed_rows = processed_rows[:head]
    elif tail is not None and tail > 0:
        processed_rows = processed_rows[-tail:]
    else:
        start = max(0, offset)
        if limit is not None and limit >= 0:
            processed_rows = processed_rows[start : start + limit]
        elif start > 0:
            processed_rows = processed_rows[start:]

    return actual_sheet, final_headers, processed_rows, total_matched_rows


def format_markdown(headers: list[str], rows: list[list[str]]) -> str:
    """转换为精简对齐的 Markdown 表格格式。"""
    if not headers and not rows:
        return "*（无数据）*"

    # 处理单元格内容：换行符转为 <br>，竖线转为 \|，避免破坏 Markdown 表格
    def sanitize(cell: str) -> str:
        s = cell.replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")
        return s.replace("|", r"\|")

    clean_headers = [sanitize(h) for h in headers]
    clean_rows = [[sanitize(c) for c in r] for r in rows]

    col_widths = [len(h) for h in clean_headers]
    for r in clean_rows:
        for i, c in enumerate(r):
            if i < len(col_widths):
                col_widths[i] = max(col_widths[i], len(c))
            else:
                col_widths.append(len(c))

    # 最小宽度为 3
    col_widths = [max(w, 3) for w in col_widths]

    lines = []
    # 表头
    header_line = "| " + " | ".join(h.ljust(col_widths[i]) for i, h in enumerate(clean_headers)) + " |"
    sep_line = "| " + " | ".join("-" * col_widths[i] for i in range(len(col_widths))) + " |"
    lines.append(header_line)
    lines.append(sep_line)

    # 数据行
    for r in clean_rows:
        row_line = "| " + " | ".join(c.ljust(col_widths[i]) for i, c in enumerate(r)) + " |"
        lines.append(row_line)

    return "\n".join(lines)


def format_tsv(headers: list[str], rows: list[list[str]], include_header: bool = True) -> str:
    """转换为 TSV 格式。"""
    output = io.StringIO()
    writer = csv.writer(output, delimiter="\t", lineterminator="\n")
    if include_header and headers:
        writer.writerow(headers)
    for r in rows:
        # 移除或转义内部 \t
        clean_r = [c.replace("\t", " ") for c in r]
        writer.writerow(clean_r)
    return output.getvalue().rstrip("\n")


def format_csv(headers: list[str], rows: list[list[str]], include_header: bool = True) -> str:
    """转换为标准 CSV 格式。"""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=",", quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
    if include_header and headers:
        writer.writerow(headers)
    for r in rows:
        writer.writerow(r)
    return output.getvalue().rstrip("\n")


def format_json(
    headers: list[str],
    rows: list[list[str]],
    mode: str = "records",
    indent: int = 2,
) -> str:
    """转换为 JSON 格式（records 列表或 rows 数组）。"""
    if mode == "rows":
        payload = {
            "headers": headers,
            "rows": rows,
        }
    else:  # records
        records = []
        for r in rows:
            rec = {}
            for i, h in enumerate(headers):
                rec[h] = r[i] if i < len(r) else ""
            records.append(rec)
        payload = records

    return json.dumps(payload, ensure_ascii=False, indent=indent)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Excel (.xlsx) 文件解析与预览工具，支持 Markdown/TSV/CSV/JSON 输出、行列切片及关键词搜索。",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("file", nargs="?", help="要解析的 .xlsx 文件路径")
    parser.add_argument(
        "-l", "--list-sheets",
        action="store_true",
        help="列出工作簿中包含的所有工作表名称及行列规模",
    )
    parser.add_argument(
        "-s", "--sheet",
        default=None,
        help="指定工作表名称或 0-based 索引（默认读取首个工作表）",
    )
    parser.add_argument(
        "-f", "--format",
        choices=["markdown", "md", "tsv", "csv", "json"],
        default="markdown",
        help="输出格式：markdown (md) / tsv / csv / json（默认: markdown）",
    )
    parser.add_argument(
        "--json-mode",
        choices=["records", "rows"],
        default="records",
        help="JSON 输出形式：records ([{col: val}]) 或 rows ({headers: [], rows: [[]]})（默认: records）",
    )
    parser.add_argument(
        "-c", "--cols",
        default=None,
        help="仅提取指定列（逗号分隔，支持列名、列字母如 A,B,D 或 1-based 数字如 1,2,4）",
    )
    parser.add_argument(
        "-q", "--search",
        default=None,
        help="按关键词筛选行（在选定列中包含指定字符串的行）",
    )
    parser.add_argument(
        "--match-case",
        action="store_true",
        help="搜索时区分大小写（默认不区分）",
    )
    parser.add_argument(
        "-n", "--head",
        type=int,
        default=None,
        help="仅显示前 N 行数据",
    )
    parser.add_argument(
        "--tail",
        type=int,
        default=None,
        help="仅显示后 N 行数据",
    )
    parser.add_argument(
        "--offset",
        type=int,
        default=0,
        help="跳过前 N 行数据（0-based）",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="限制输出的数据行数",
    )
    parser.add_argument(
        "--header-row",
        type=int,
        default=1,
        help="表头所在行号（1-based，默认 1；传入 0 表示无表头）",
    )
    parser.add_argument(
        "-o", "--out",
        default=None,
        help="输出到目标文件，默认打印到控制台 (stdout)",
    )
    parser.add_argument(
        "--info",
        action="store_true",
        help="打印总行数、列数等统计摘要信息",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    _ensure_utf8_stdout()
    parser = build_parser()
    args = parser.parse_args(argv)

    if not args.file:
        parser.print_help()
        return 1

    file_path = Path(args.file)
    if not file_path.exists():
        sys.stderr.write(f"错误: 文件不存在: {file_path}\n")
        return 1

    # 1. 仅列出 Sheet
    if args.list_sheets:
        try:
            sheets = list_sheets(file_path)
            print(f"文件: {file_path.name}")
            print(f"工作表总数: {len(sheets)}")
            print("-" * 50)
            for s in sheets:
                row_str = f"{s['max_row']} 行" if s['max_row'] is not None else "未知行"
                col_str = f"{s['max_column']} 列" if s['max_column'] is not None else "未知列"
                print(f"  [{s['index']}] {s['name']} ({row_str}, {col_str})")
            return 0
        except Exception as e:
            sys.stderr.write(f"读取工作表列表失败: {e}\n")
            return 1

    # 2. 读取并过滤数据
    try:
        actual_sheet, headers, rows, total_matched = load_sheet_data(
            file_path=file_path,
            sheet=args.sheet,
            header_row=args.header_row,
            cols=args.cols,
            search=args.search,
            search_ignore_case=not args.match_case,
            offset=args.offset,
            limit=args.limit,
            head=args.head,
            tail=args.tail,
        )
    except Exception as e:
        sys.stderr.write(f"读取 Excel 失败: {e}\n")
        return 1

    # 统计信息
    info_header = (
        f"# 工作表: {actual_sheet} | 总匹配行数: {total_matched} | 当前显示: {len(rows)} 行 | 列数: {len(headers)}\n"
    )

    fmt = args.format.lower()
    if fmt in ("markdown", "md"):
        content = format_markdown(headers, rows)
    elif fmt == "tsv":
        content = format_tsv(headers, rows, include_header=args.header_row > 0)
    elif fmt == "csv":
        content = format_csv(headers, rows, include_header=args.header_row > 0)
    elif fmt == "json":
        content = format_json(headers, rows, mode=args.json_mode)
    else:
        content = format_markdown(headers, rows)

    # 写入文件或打印输出
    if args.out:
        out_path = Path(args.out)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(content, encoding="utf-8")
        print(f"成功导出 {len(rows)} 行数据到: {out_path} (工作表: {actual_sheet})")
    else:
        if args.info:
            print(info_header)
        print(content)

    return 0


if __name__ == "__main__":
    sys.exit(main())
